
from PIL import ImageMode
from PIL import ImageMode
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
from tkcalendar import Calendar
from docx import Document
from docx.shared import Mm, Pt
import sqlite3
import os
import subprocess
import sys

def convert_pdf_to_docx(pdf_path, docx_path=None):
    """Converts a generated PDF letterhead/prescription file into an editable MS Word (.docx) document."""
    if not docx_path:
        docx_path = os.path.splitext(pdf_path)[0] + ".docx"

    target_path = docx_path
    counter = 1
    while True:
        try:
            if os.path.exists(target_path):
                with open(target_path, "a"):
                    pass
            break
        except (PermissionError, IOError):
            base, ext = os.path.splitext(docx_path)
            target_path = f"{base}_{counter}{ext}"
            counter += 1

    try:
        from pdf2docx import Converter
        cv = Converter(pdf_path)
        cv.convert(target_path)
        cv.close()
        return target_path
    except Exception as exc:
        print(f"pdf2docx conversion error: {exc}")
        return target_path

def open_file(filepath):
    """Opens a PDF or DOCX file using system default application."""
    try:
        if hasattr(os, "startfile"):
            os.startfile(filepath)
        elif sys.platform.startswith("linux"):
            subprocess.Popen(["xdg-open", filepath])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", filepath])
        else:
            subprocess.Popen(["cmd", "/c", "start", "", filepath])
    except Exception as exc:
        messagebox.showwarning("Open File", f"Could not open file automatically.\n{exc}")

class Letter:
    def __init__(self, app, mode="letter"):
        self.app = app
        self.init_db()
        if mode == "letter":
            self.Letter_Pad()
        elif mode == "plain":
            self.Plain_Priscription()
        elif mode == "form":
            self.todayapp()
    
    def init_db(self):
        self.conn = sqlite3.connect("dental.db")
        self.cursor = self.conn.cursor()
        self.cursor.execute("CREATE TABLE IF NOT EXISTS letterpad (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, address TEXT, phone TEXT, email TEXT, website TEXT, logo TEXT)")
        self.conn.commit()

    def get_doctors_from_db(self):
        default_consultants = [
            {"id": -1, "name": "Dr.ANOOP KUMAR. B D S", "role": "DENTAL SURGEON", "reg": "Reg.No. 1/32"},
        ]
        default_visiting = [
            {"id": -2, "name": "Dr.JUSTIN MATHEW. MDS",       "role": "Oral & Maxillo Facial Surgeon", "reg": "Reg.No. 5492"},
            {"id": -3, "name": "Dr.KRISHNA KUMAR. MDS",        "role": "Paedodontist",                  "reg": "Reg.No. 7729"},
            {"id": -4, "name": "Dr.JOSEPH J PULIKKOTTIL. MDS", "role": "Periodontist & Implantologist", "reg": "Reg.No. 2418"},
            {"id": -5, "name": "Dr.TERRY THOMAS. MDS",         "role": "Orthodontist",                  "reg": "Reg.No. 4807"},
            {"id": -6, "name": "Dr.SHIBHU SREEDHAR. MDS",      "role": "Endodontist",                   "reg": "Reg.No. 15619-A"},
            {"id": -7, "name": "Dr.RENJITH RAJ. MDS",          "role": "Endodontist",                   "reg": "Reg.No. 8171"},
            {"id": -8, "name": "Dr.SIJO P MATHEW. MDS",        "role": "Endodontist",                   "reg": "Reg.No. 8982"},
        ]
        try:
            db_path = getattr(self.app, "db_path", "dental.db")
            if not os.path.exists(db_path):
                script_dir = os.path.dirname(os.path.abspath(__file__))
                db_path = os.path.join(script_dir, "dental.db")
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute("PRAGMA table_info(Doctors)")
            cols = [c[1] for c in cursor.fetchall()]
            if "Is_Consultant" not in cols or "Is_Visiting" not in cols:
                conn.close()
                return default_consultants, default_visiting

            cursor.execute("SELECT id, Reg_No, First_Name, Last_Name, Qualification, Designation, Is_Consultant, Is_Visiting FROM Doctors")
            rows = cursor.fetchall()
            conn.close()

            consultants = []
            visiting = []
            for doc_id, reg, first, last, qual, desig, is_cons, is_vis in rows:
                fname = f"{first or ''} {last or ''}".strip()
                full_name = f"Dr. {fname}".strip() if fname else "Dr."
                if qual:
                    full_name += f". {qual}"
                doc_dict = {
                    "id": doc_id,
                    "name": full_name,
                    "role": desig or "",
                    "reg": f"Reg.No. {reg}" if reg else ""
                }
                if is_cons:
                    consultants.append(doc_dict)
                if is_vis:
                    visiting.append(doc_dict)

            if not consultants and not visiting:
                return default_consultants, default_visiting

            if not consultants:
                consultants = default_consultants
            if not visiting:
                visiting = default_visiting

            return consultants, visiting
        except Exception:
            return default_consultants, default_visiting

    def get_clinic_phone_settings(self):
        default_mobile = "9446046868"
        default_work = "216858"
        try:
            db_path = getattr(self.app, "db_path", "dental.db")
            if not os.path.exists(db_path):
                script_dir = os.path.dirname(os.path.abspath(__file__))
                db_path = os.path.join(script_dir, "dental.db")
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT key, value FROM Clinic_Settings WHERE key IN ('clinic_mobile', 'clinic_work', 'clinic_phone', 'clinic_resi')")
            rows = dict(cursor.fetchall())
            conn.close()

            mobile = rows.get("clinic_mobile") or rows.get("clinic_phone") or default_mobile
            work = rows.get("clinic_work") or rows.get("clinic_resi") or default_work

            phone_str = mobile if ("Clinic" in mobile or "clinic" in mobile) else f"Clinic : {mobile}"
            resi_str = work if ("Resi" in work or "resi" in work) else f"Resi   : {work}"
            return phone_str, resi_str
        except Exception:
            return f"Clinic : {default_mobile}", f"Resi   : {default_work}"

    def get_clinic_hours_setting(self):
        default_hours = "Clinic Hours : 10:00 AM to 07:00 PM,  Tuesday Holiday"
        try:
            db_path = getattr(self.app, "db_path", "dental.db")
            if not os.path.exists(db_path):
                script_dir = os.path.dirname(os.path.abspath(__file__))
                db_path = os.path.join(script_dir, "dental.db")
            conn = sqlite3.connect(db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT key, value FROM Clinic_Settings WHERE key IN ('clinic_open_time', 'clinic_close_time', 'clinic_holiday', 'clinic_hours')")
            rows = dict(cursor.fetchall())
            conn.close()

            if "clinic_hours" in rows and rows["clinic_hours"]:
                return rows["clinic_hours"]

            open_time = rows.get("clinic_open_time", "10:00 AM")
            close_time = rows.get("clinic_close_time", "07:00 PM")
            holiday = rows.get("clinic_holiday", "Tuesday Holiday")

            if holiday:
                holiday_str = holiday if ("Holiday" in holiday or "holiday" in holiday) else f"{holiday} Holiday"
                return f"Clinic Hours : {open_time} to {close_time},  {holiday_str}"
            else:
                return f"Clinic Hours : {open_time} to {close_time}"
        except Exception:
            return default_hours
    

    def Letter_Pad(self):
        # ---- ReportLab import ----
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm as rl_mm
            from reportlab.pdfgen import canvas as rl_canvas
            from reportlab.lib.colors import HexColor
            from reportlab.lib.utils import ImageReader
            reportlab_ok = True
            PAGE_W, PAGE_H = A4          # points  (595 × 842)
        except ImportError:
            reportlab_ok = False
            PAGE_W, PAGE_H = 595.0, 842.0
            rl_mm = 2.8346
            HexColor = None
            ImageReader = None
            rl_canvas = None

        # ── Paths & clinic data ──────────────────────────────────────────
        SCRIPT_DIR     = os.path.dirname(os.path.abspath(__file__))
        LOGO_PATH      = os.path.join(SCRIPT_DIR, "Dental_logo.png")
        CLINIC_NAME    = "ANUPAM DENTAL CLINIC"
        CLINIC_ADDRESS = "West Gate Vaikom - 686141"
        CLINIC_PHONE, CLINIC_RESI = self.get_clinic_phone_settings()
        CLINIC_HOURS   = self.get_clinic_hours_setting()

        db_consultants, db_visiting = self.get_doctors_from_db()
        CONSULTANTS = list(db_consultants)
        VISITING_DOCTORS = list(db_visiting)

        # ==============================================================
        # PDF generation (ReportLab)
        # ==============================================================
        def draw_logo(c, cx: float, cy: float, r: float):
            if os.path.isfile(LOGO_PATH) and ImageReader:
                img = ImageReader(LOGO_PATH)
                iw, ih = img.getSize()
                box = 2 * r
                scale = min(box / iw, box / ih)
                w, h = iw * scale, ih * scale
                c.drawImage(img, cx - w / 2, cy - h / 2, width=w, height=h, mask="auto", preserveAspectRatio=True)
                return
            
            # Fallback
            c.saveState()
            c.setDash(1, 2)
            c.setLineWidth(1)
            c.circle(cx, cy, r, stroke=1, fill=0)
            c.circle(cx, cy, r - 2 * rl_mm, stroke=1, fill=0)
            c.restoreState()
            c.saveState()
            c.translate(cx, cy)
            c.rotate(90)
            c.setFont("Helvetica-Bold", 5.5)
            c.drawCentredString(0, r - 5 * rl_mm, "ANUPAM")
            c.restoreState()
            c.setFont("Helvetica", 4.5)
            c.drawCentredString(cx, cy - 1.5 * rl_mm, "DENTAL")
            c.drawCentredString(cx, cy - 4.5 * rl_mm, "CLINIC")

        def draw_letterhead(c):
            if HexColor:
                ink = HexColor("#1a1a1a")
                c.setFillColor(ink)
                c.setStrokeColor(ink)
            
            logo_cx, logo_cy, logo_r = 30 * rl_mm, PAGE_H - 20 * rl_mm, 15 * rl_mm
            draw_logo(c, logo_cx, logo_cy, logo_r)
            
            header_cx = PAGE_W / 2
            c.setFont("Times-Bold", 20)
            c.drawCentredString(header_cx, PAGE_H - 15 * rl_mm, CLINIC_NAME)
            c.setFont("Helvetica", 8.5)
            c.drawCentredString(header_cx, PAGE_H - 22 * rl_mm, CLINIC_ADDRESS)
            c.setFont("Helvetica", 9)
            c.drawCentredString(PAGE_W - 28 * rl_mm, PAGE_H - 30 * rl_mm, f"{CLINIC_PHONE}")
            c.drawCentredString(PAGE_W - 32 * rl_mm, PAGE_H - 37 * rl_mm, f"{CLINIC_RESI}")
            
            top_line_y = PAGE_H - 47 * rl_mm
            margin = 10 * rl_mm
            right_x = PAGE_W - 2 * rl_mm
            c.setLineWidth(0.8)
            c.line(margin, top_line_y, right_x, top_line_y)
            
            divider_x = margin + 70 * rl_mm
            bottom_line_y = 12 * rl_mm
            c.setLineWidth(0.8)
            c.line(divider_x, top_line_y, divider_x, bottom_line_y)
            c.line(divider_x, top_line_y, right_x, top_line_y)
            c.line(right_x, top_line_y, right_x, bottom_line_y)
            c.line(divider_x, bottom_line_y, right_x, bottom_line_y)
            
            left_x = margin + 1 * rl_mm
            y = top_line_y - 8 * rl_mm
            
            def section_label(text, size=9, gap=5.2):
                nonlocal y
                c.setFont("Helvetica-Bold", size)
                c.drawString(left_x, y, text)
                y -= gap * rl_mm
            
            def doctor_entry(doc, gap=4.6, size=7.8):
                nonlocal y
                c.setFont("Helvetica", size)
                for line in (doc["name"], doc["role"], doc["reg"]):
                    c.drawString(left_x, y, line)
                    y -= gap * rl_mm
                y -= 1.5 * rl_mm
            
            section_label("Consultants :")
            for doc in CONSULTANTS:
                doctor_entry(doc)
            
            y -= 6 * rl_mm
            section_label("Visiting :")
            for doc in VISITING_DOCTORS:
                doctor_entry(doc)
            
            c.setFont("Helvetica", 8.5)
            c.drawCentredString(header_cx, 6 * rl_mm, CLINIC_HOURS)

        def generate_pdf(filepath: str):
            if not rl_canvas: return
            c = rl_canvas.Canvas(filepath, pagesize=(PAGE_W, PAGE_H))
            draw_letterhead(c)
            c.showPage()
            c.save()

        # ==============================================================
        # Build Tkinter workspace
        # ==============================================================
        self.app.clear_workspace()
        self.app.workspace = tk.Frame(self.app.root, bd=3, relief="solid")
        self.app.workspace.pack(padx=10, pady=10, fill="both", expand=True)
        ws = self.app.workspace

        AVAIL_H = 500
        AVAIL_W = 500
        base_scale = min(AVAIL_H / PAGE_H, AVAIL_W / PAGE_W)
        self.preview_scale = base_scale
        CW      = int(PAGE_W * self.preview_scale)
        CH      = int(PAGE_H * self.preview_scale)

        outer = tk.Frame(ws, bg="#c0c0c0")
        outer.pack(fill="both", expand=True, padx=10, pady=(6, 0))

        vsb = tk.Scrollbar(outer, orient="vertical")
        vsb.pack(side="right", fill="y")
        hsb = tk.Scrollbar(outer, orient="horizontal")
        hsb.pack(side="bottom", fill="x")

        cv = tk.Canvas(outer, bg="#c0c0c0", width=CW+10, height=CH+1,
                       scrollregion=(0, 0, CW + 20, CH + 20),
                       yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        cv.pack(fill="both", expand=True)
        vsb.config(command=cv.yview)
        hsb.config(command=cv.xview)
        
        state = {
            "current_page": 0,
            "pages": [0]
        }

        # Preview Controls (Zoom & Navigation)
        ctrl_frame = tk.Frame(ws)
        ctrl_frame.pack(fill="x", pady=4)

        def prev_page():
            if state["current_page"] > 0:
                draw_page_preview(state["current_page"] - 1)

        def next_page():
            if state["current_page"] < len(state["pages"]) - 1:
                draw_page_preview(state["current_page"] + 1)

        # Page buttons hidden for single-page preview

        def zoom_in():
            if self.preview_scale < 2.0:
                self.preview_scale = round(min(2.0, self.preview_scale + 0.1), 1)
                draw_page_preview(state["current_page"])

        def zoom_out():
            if self.preview_scale > 0.25:
                self.preview_scale = round(max(0.25, self.preview_scale - 0.1), 1)
                draw_page_preview(state["current_page"])

        def on_mouse_wheel(event):
            if getattr(event, "state", 0) & 1:  # Shift pressed -> Horizontal scroll
                if getattr(event, "delta", 0) > 0 or getattr(event, "num", 0) == 4:
                    cv.xview_scroll(-1, "units")
                else:
                    cv.xview_scroll(1, "units")
            else:
                if getattr(event, "delta", 0) > 0 or getattr(event, "num", 0) == 4:
                    zoom_in()
                else:
                    zoom_out()
            return "break"

        cv.bind("<MouseWheel>", on_mouse_wheel)
        cv.bind("<Shift-MouseWheel>", on_mouse_wheel)
        cv.bind("<Button-4>", on_mouse_wheel)
        cv.bind("<Button-5>", on_mouse_wheel)

        btn_zoom_out = tk.Button(ctrl_frame, text="-", font=("Arial", 9, "bold"), command=zoom_out)
        btn_zoom_out.pack(side="left", padx=5)

        zoom_lbl = tk.Label(ctrl_frame, text="100%", font=("Arial", 10), bg="white")
        zoom_lbl.pack(side="left", padx=5)

        btn_zoom_in = tk.Button(ctrl_frame, text="+", font=("Arial", 9, "bold"), command=zoom_in)
        btn_zoom_in.pack(side="left", padx=5)

        def draw_page_preview(page_idx):
            state["current_page"] = page_idx
            cv.delete("all")
            
            scale = self.preview_scale
            w_scaled = int(PAGE_W * scale)
            h_scaled = int(PAGE_H * scale)
            mm_px = scale * 2.8346

            cv.update_idletasks()
            win_w = cv.winfo_width()
            content_w = max(win_w, w_scaled + 80) if win_w > 50 else w_scaled + 80
            OX = max(40, (content_w - w_scaled) // 2)
            OY = 10
            total_w = OX * 2 + w_scaled
            total_h = h_scaled + 40
            
            cv.config(scrollregion=(0, 0, total_w, total_h))
            
            cv.create_rectangle(OX+4, OY+4, OX+w_scaled+4, OY+h_scaled+4, fill="#888888", outline="")
            cv.create_rectangle(OX, OY, OX+w_scaled, OY+h_scaled, fill="white", outline="#aaaaaa", width=1)

            def ppx(pt):  return OX + int(pt * scale)
            def ppy(pt):  return OY + int((PAGE_H - pt) * scale)
            def pcx():    return OX + w_scaled // 2
            def smm(v):   return int(v * mm_px)

            zoom_lbl.config(text=f"{int(scale / base_scale * 100)}%")

            logo_cx = ppx(30*rl_mm)
            logo_cy = ppy(PAGE_H - 20*rl_mm)
            logo_r  = smm(15)
            if os.path.isfile(LOGO_PATH):
                try:
                    from PIL import Image as _PI, ImageTk as _ITk
                    _img = _PI.open(LOGO_PATH)
                    _img.thumbnail((logo_r*2, logo_r*2))
                    self._lp_logo = _ITk.PhotoImage(_img)
                    cv.create_image(logo_cx, logo_cy, image=self._lp_logo)
                except Exception:
                    cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                                   logo_cx+logo_r, logo_cy+logo_r, outline="#555")
            else:
                cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                               logo_cx+logo_r, logo_cy+logo_r, outline="#555")

            cv.create_text(pcx(), ppy(PAGE_H - 15*rl_mm), text=CLINIC_NAME,
                           font=("Times New Roman", max(10, smm(6)), "bold"), fill="#1a1a1a")
            cv.create_text(pcx(), ppy(PAGE_H - 22*rl_mm), text=CLINIC_ADDRESS,
                           font=("Arial", max(8, smm(3))), fill="#333")
            cv.create_text(ppx(PAGE_W - 28*rl_mm), ppy(PAGE_H - 30*rl_mm),text=f"{CLINIC_PHONE}",
                           font=("Arial", max(8, smm(3))), fill="#333")
            cv.create_text(ppx(PAGE_W - 32*rl_mm), ppy(PAGE_H - 37*rl_mm),text=f"{CLINIC_RESI}",
                           font=("Arial", max(8, smm(3))), fill="#333")

            rule_y   = ppy(PAGE_H - 47*rl_mm)
            marg_px  = ppx(10*rl_mm)
            right_px = OX + w_scaled - smm(2)
            cv.create_line(marg_px, rule_y, right_px, rule_y, fill="#333", width=1)

            div_x = marg_px + smm(70)
            bot_y = ppy(12*rl_mm)
            cv.create_line(div_x, rule_y, div_x, bot_y, fill="#555", width=1)
            cv.create_rectangle(div_x, rule_y, right_px, bot_y, outline="#555", width=1)

            lx  = marg_px + smm(1)
            ly  = rule_y + smm(8)
            lh  = smm(4.6)
            lhs = smm(5.2)

            cv.create_text(lx, ly, text="Consultants :",
                           font=("Arial", max(7, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            ly += lhs
            for doc in CONSULTANTS:
                cv.create_text(lx+smm(3), ly, text=doc["name"],
                               font=("Arial", max(6, smm(2.5))), fill="#1a1a1a", anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=doc["role"],
                               font=("Arial", max(5, smm(2))), fill="#333",    anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=doc["reg"],
                               font=("Arial", max(5, smm(2))), fill="#555",    anchor="nw"); ly += lh + smm(1.5)

            ly += smm(6)
            cv.create_text(lx, ly, text="Visiting :",
                           font=("Arial", max(7, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            ly += lhs
            for doc in VISITING_DOCTORS:
                cv.create_text(lx+smm(3), ly, text=doc["name"],
                               font=("Arial", max(6, smm(2.5))), fill="#1a1a1a", anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=doc["role"],
                               font=("Arial", max(5, smm(2))), fill="#333",    anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=doc["reg"],
                               font=("Arial", max(5, smm(2))), fill="#555",    anchor="nw"); ly += lh + smm(1.5)

            cv.create_text(pcx(), ppy(6*rl_mm), text=CLINIC_HOURS,
                           font=("Arial", max(6, smm(2.5))), fill="#555")

        draw_page_preview(0)

        btn_bar = tk.Frame(ws)
        btn_bar.pack(pady=8)

        def _open_pdf(path):
            open_file(path)

        def _generate_pdf():
            if not reportlab_ok:
                messagebox.showerror("Missing Library",
                    "ReportLab is required.\nRun:  pip install reportlab")
                return
            filepath = os.path.join(SCRIPT_DIR, "Anupam_Dental_Letterhead.pdf")
            try:
                generate_pdf(filepath)
                docx_path = convert_pdf_to_docx(filepath)
                messagebox.showinfo("Done", f"Letterhead saved successfully:\n📄 PDF: {filepath}\n📝 Word: {docx_path}")
                _open_pdf(filepath)
            except Exception as exc:
                messagebox.showerror("Error", f"PDF generation failed:\n{exc}")

        def _print_now():
            if not reportlab_ok:
                messagebox.showerror("Missing Library",
                    "ReportLab is required.\nRun:  pip install reportlab")
                return
            tmp = os.path.join(SCRIPT_DIR, "_letterhead_temp.pdf")
            try:
                generate_pdf(tmp)
                convert_pdf_to_docx(tmp)
                _open_pdf(tmp)
            except Exception as exc:
                messagebox.showerror("Error", f"Could not open PDF:\n{exc}")
        
        def _open_word(out_path, clinic_name, clinic_address, clinic_phone, clinic_resi, clinic_hours):
            """Generate a .docx letterhead and open it."""
            from docx import Document as _Doc
            from docx.shared import Pt, Mm, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            import copy

            doc = _Doc()

            # ── Page margins ──────────────────────────────────────────────
            section = doc.sections[0]
            section.page_width  = Mm(210)
            section.page_height = Mm(297)
            section.top_margin    = Mm(10)
            section.bottom_margin = Mm(12)
            section.left_margin   = Mm(14)
            section.right_margin  = Mm(14)

            def add_centered(text, bold=False, size=11, color=None, space_after=0):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after  = Pt(space_after)
                p.paragraph_format.space_before = Pt(0)
                run = p.add_run(text)
                run.bold = bold
                run.font.size = Pt(size)
                if color:
                    run.font.color.rgb = RGBColor(*color)
                return p

            def add_left(text, bold=False, size=9, indent_mm=0, space_after=0):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after  = Pt(space_after)
                p.paragraph_format.space_before = Pt(0)
                if indent_mm:
                    p.paragraph_format.left_indent = Mm(indent_mm)
                run = p.add_run(text)
                run.bold = bold
                run.font.size = Pt(size)
                return p

            def add_rule():
                """Insert a thin horizontal border below the previous paragraph."""
                p = doc.add_paragraph()
                p.paragraph_format.space_after  = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                pPr = p._p.get_or_add_pPr()
                pBdr = OxmlElement('w:pBdr')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'),   'single')
                bottom.set(qn('w:sz'),    '6')
                bottom.set(qn('w:space'), '1')
                bottom.set(qn('w:color'), '333333')
                pBdr.append(bottom)
                pPr.append(pBdr)

            # ── Clinic header ─────────────────────────────────────────────
            table = doc.add_table(rows=1, cols=3)
            table.autofit = False
            for cell in table.columns[0].cells: cell.width = Mm(40)
            for cell in table.columns[1].cells: cell.width = Mm(102)
            for cell in table.columns[2].cells: cell.width = Mm(440)
            # for cell in table.columns[3].cells: cell.width = Mm(440)

            # Logo in the left cell
            cell_logo = table.cell(0, 0)
            if os.path.isfile(LOGO_PATH):
                p_logo = cell_logo.paragraphs[0]
                p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_logo.add_run()
                run.add_picture(LOGO_PATH, width=Mm(25))

            # Clinic details in the middle cell
            cell_mid = table.cell(0, 1)
            def add_cell_centered(cell, text, bold=False, size=11, color=None, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER):
                if len(cell.paragraphs) == 0:
                    p = cell.add_paragraph()
                else:
                    p = cell.paragraphs[0]
                p.alignment = align
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(space_after)
                run = p.add_run(text)
                run.bold = bold
                run.font.size = Pt(size)
                if color:
                    run.font.color.rgb = RGBColor(*color)

            add_cell_centered(cell_mid, clinic_name,    bold=True,  size=18, space_after=2)
            add_cell_centered(cell_mid, clinic_address, bold=False, size=10, space_after=1)
            
            # Clinic details in the right cell
            cell_right = table.cell(0, 2)
            add_cell_centered(cell_right, clinic_phone, bold=False, size=9, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
            add_cell_centered(cell_right, clinic_resi,  bold=False, size=9, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)

            add_rule()  

            # ── Consultants ───────────────────────────────────────────────
            add_left("Consultants :", bold=True, size=9, space_after=0)
            for doc_info in CONSULTANTS:
                add_left(doc_info["name"], bold=False, size=8, indent_mm=5, space_after=0)
                add_left(doc_info["role"], bold=False, size=6, indent_mm=5, space_after=0)
                add_left(doc_info["reg"],  bold=False, size=6, indent_mm=5, space_after=3)

            add_left("Visiting :", bold=True, size=9, space_after=0)
            for doc_info in VISITING_DOCTORS:
                add_left(doc_info["name"], bold=False, size=8, indent_mm=5, space_after=0)
                add_left(doc_info["role"], bold=False, size=6, indent_mm=5, space_after=0)
                add_left(doc_info["reg"],  bold=False, size=6, indent_mm=5, space_after=3)

            add_rule()

            # ── Footer ────────────────────────────────────────────────────
            add_centered(clinic_hours, bold=False, size=8, space_after=0)

            # ── Save & open ───────────────────────────────────────────────
            doc.save(out_path)
            try:
                if hasattr(os, "startfile"):
                    os.startfile(out_path)
                elif sys.platform.startswith("linux"):
                    subprocess.Popen(["xdg-open", out_path])
                elif sys.platform == "darwin":
                    subprocess.Popen(["open", out_path])
                else:
                    subprocess.Popen(["cmd", "/c", "start", "", out_path])
            except Exception as exc:
                messagebox.showwarning("Open Word", f"Could not open the document automatically.\n{exc}")

        def _generate_word():
            try:
                from docx import Document as _check  # noqa – just verify installed
            except ImportError:
                messagebox.showerror("Missing Library",
                    "python-docx is required.\nRun:  pip install python-docx")
                return
            filepath = os.path.join(SCRIPT_DIR, "Anupam_Dental_Letterhead.docx")
            try:
                tmp_pdf = os.path.join(SCRIPT_DIR, "_letterhead_temp.pdf")
                generate_pdf(tmp_pdf)
                docx_path = convert_pdf_to_docx(tmp_pdf, filepath)
                messagebox.showinfo("Done", f"Word document saved:\n{docx_path}")
                open_file(docx_path)
            except Exception as exc:
                messagebox.showerror("Error", f"Word generation failed:\n{exc}")

        tk.Button(btn_bar, text="📄  Generate PDF", font=("Arial", 11), width=16,
                  bg="#1565C0", fg="white", command=_generate_pdf).grid(row=0, column=0, padx=8)
        tk.Button(btn_bar, text="🖨  Open / Print", font=("Arial", 11), width=16,
                  bg="#2E7D32", fg="white", command=_print_now).grid(row=0, column=1, padx=8)
        tk.Button(btn_bar, text="📝  Word", font=("Arial", 11), width=12,
                  bg="#6A1B9A", fg="white", command=_generate_word).grid(row=0, column=2, padx=8)
        tk.Button(btn_bar, text="Close",            font=("Arial", 11), width=10,
                  command=self.close).grid(row=0, column=3, padx=8)

    def Plain_Priscription(self):
        
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import HexColor
        from reportlab.lib.utils import ImageReader
        
        # Half A4, portrait orientation: same half-sheet as the landscape version,
        # just rotated 90 degrees (148.5 x 210 mm - i.e. A5 portrait)
        A4_W, A4_H = A4
        PAGE_W, PAGE_H = A4_H / 2, A4_W
        
        # ----------------------------------------------------------------------
        # Logo image - place "Dental_logo.png" in the same folder as this script
        # (or change LOGO_PATH to point wherever the file lives). If the file
        # can't be found, a simple drawn placeholder circle is used instead.
        # ----------------------------------------------------------------------
        SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
        LOGO_PATH = os.path.join(SCRIPT_DIR, "Dental_logo.png")
        
        # ----------------------------------------------------------------------
        # Editable clinic data - change these to update the letterhead content
        # ----------------------------------------------------------------------
        CLINIC_NAME = "ANUPAM DENTAL CLINIC"
        CLINIC_ADDRESS = "West Gate Vaikom - 686141"
        CLINIC_PHONE, CLINIC_RESI = self.get_clinic_phone_settings()
        CLINIC_HOURS = self.get_clinic_hours_setting()
         
        db_consultants, db_visiting = self.get_doctors_from_db()
        CONSULTANTS = list(db_consultants)
        VISITING_DOCTORS = list(db_visiting)
        
        
        # ----------------------------------------------------------------------
        # PDF drawing
        # ----------------------------------------------------------------------
        def draw_logo(c: canvas.Canvas, cx: float, cy: float, r: float):
            if os.path.isfile(LOGO_PATH):
                img = ImageReader(LOGO_PATH)
                iw, ih = img.getSize()
                box = 2 * r
                scale = min(box / iw, box / ih)
                w, h = iw * scale, ih * scale
                c.drawImage(
                    img,
                    cx - w / 2,
                    cy - h / 2,
                    width=w,
                    height=h,
                    mask="auto",  # respects transparency (PNG alpha channel)
                    preserveAspectRatio=True,
                )
                return
        
            # ---- Fallback: simple drawn placeholder seal (used if image missing) ----
            c.saveState()
            c.setDash(1, 2)
            c.setLineWidth(1)
            c.circle(cx, cy, r, stroke=1, fill=0)
            c.circle(cx, cy, r - 2 * mm, stroke=1, fill=0)
            c.restoreState()
        
            c.saveState()
            c.translate(cx, cy)
            c.rotate(90)
            c.setFont("Helvetica-Bold", 5.5)
            c.drawCentredString(0, r - 5 * mm, "ANUPAM")
            c.restoreState()
        
            c.setFont("Helvetica", 4.5)
            c.drawCentredString(cx, cy - 1.5 * mm, "DENTAL")
            c.drawCentredString(cx, cy - 4.5 * mm, "CLINIC")
        
        
        def draw_letterhead(c: canvas.Canvas):
            """Draws the full letterhead on the given ReportLab canvas (portrait half-A4 page)."""
            ink = HexColor("#1a1a1a")
            c.setFillColor(ink)
            c.setStrokeColor(ink)
        
            # ---------- Logo (top-left) ----------
            logo_cx, logo_cy, logo_r = 20 * mm, PAGE_H - 16 * mm, 9 * mm
            draw_logo(c, logo_cx, logo_cy, logo_r)
        
            # ---------- Header text (stacked: title / address / phones) ----------
            header_cx = PAGE_W * 0.60 
            header_cx1 = PAGE_W * 0.85
        
            c.setFont("Times-Bold", 20)
            c.drawCentredString(header_cx, PAGE_H - 15 * mm, CLINIC_NAME)
        
            c.setFont("Helvetica", 8.5)
            c.drawCentredString(header_cx, PAGE_H - 20 * mm, CLINIC_ADDRESS)
        
            c.setFont("Helvetica", 9)
            c.drawCentredString(header_cx1, PAGE_H - 25 * mm, CLINIC_PHONE)

            c.setFont("Helvetica", 9)
            c.drawCentredString(header_cx1, PAGE_H - 30 * mm, CLINIC_RESI)
        
            # ---------- Horizontal rule under header ----------
            top_line_y = PAGE_H - 32 * mm
            margin = 1 * mm
            c.setLineWidth(0.8)
            c.line(margin, top_line_y, PAGE_W - margin, top_line_y) 
            c.line(margin + 80, top_line_y - 470, PAGE_W - margin, top_line_y - 470)
        
            # ---------- Divider + writing-area box (right side, like the original) ----------
            divider_x = margin + 50 * mm
            bottom_line_y = 12 * mm
            c.setLineWidth(0.8)
            c.line(divider_x, top_line_y, divider_x, bottom_line_y)
            c.line(divider_x, top_line_y, PAGE_W - margin, top_line_y)
            c.line(PAGE_W - margin, top_line_y, PAGE_W - margin, bottom_line_y)
            c.line(divider_x, bottom_line_y, PAGE_W - margin, bottom_line_y)

            #-------------Doctor's Notes--------------------------------------
            note_divider_x = divider_x + 5*mm
            note_width = (PAGE_W-margin) - note_divider_x - 5*mm
            note_height= 4*mm 
            note_y = bottom_line_y + 53*mm
            
            c.rect( note_divider_x,note_y,note_width,note_height,stroke=1,fill=0)
            c.setFont("Helvetica", 9)
            c.drawCentredString(note_divider_x + note_width / 2,note_y + 1 * mm,"Doctor's Notes")


            #-------------Date & Time--------------------------------------

            date_w = 30 * mm
            date_h = 19 * mm
            date_x = PAGE_W - margin - date_w
            date_y = bottom_line_y + 143 * mm

            # label inside box (top-left)
            c.setFont("Helvetica", 8)
            c.drawString(date_x, date_y + date_h, "Date :")
        
            # ---------- Left column: Consultants / Visiting doctors ----------
            left_x = margin + 1 * mm
            y = top_line_y - 8 * mm
        
            def section_label(text, size=9, gap=5.2):
                nonlocal y
                c.setFont("Helvetica-Bold", size)
                c.drawString(left_x, y, text)
                y -= gap * mm
        
            def doctor_entry(doc, gap=2.0, size=7.8):
                nonlocal y
                c.setFont("Helvetica", size)
                c.drawString((left_x +10), y, doc["name"])
                y -= 3.5 * mm
                c.drawString((left_x +10), y, doc["role"])
                y -= 3.2 * mm
                c.drawString((left_x +10), y, doc["reg"])
                y -= 4.5 * mm
                y-=gap * mm
            section_label("Consultants :")
            for doc in CONSULTANTS:
                doctor_entry(doc)
        
            y -= 6 * mm
            section_label("Visiting :")
            for doc in VISITING_DOCTORS:
                doctor_entry(doc)

            # ---------- Footer ----------
            c.setFont("Helvetica", 8.5)
            c.drawCentredString(header_cx - 50, 6 * mm, CLINIC_HOURS)
        
        
        def generate_pdf(filepath: str):
            """Creates the letterhead PDF (portrait half-A4 page) at the given filepath."""
            c = canvas.Canvas(filepath, pagesize=(PAGE_W, PAGE_H))
            draw_letterhead(c)
            c.showPage()
            c.save()
        
        
        # ----------------------------------------------------------------------
        # Tkinter GUI
        # ----------------------------------------------------------------------
        self.app.clear_workspace()
        self.app.workspace = tk.Frame(self.app.root, bd=3, relief="solid")
        self.app.workspace.pack(padx=10, pady=10, fill="both", expand=True)
        ws = self.app.workspace

        AVAIL_H = 500
        AVAIL_W = 500
        base_scale = min(AVAIL_H / PAGE_H, AVAIL_W / PAGE_W)
        self.preview_scale = base_scale
        CW      = int(PAGE_W * self.preview_scale)
        CH      = int(PAGE_H * self.preview_scale)

        outer = tk.Frame(ws, bg="#c0c0c0")
        outer.pack(fill="both", expand=True, padx=10, pady=(6, 0))

        vsb = tk.Scrollbar(outer, orient="vertical")
        vsb.pack(side="right", fill="y")
        hsb = tk.Scrollbar(outer, orient="horizontal")
        hsb.pack(side="bottom", fill="x")

        cv = tk.Canvas(outer, bg="#c0c0c0", width=CW+10, height=CH+1,
                       scrollregion=(0, 0, CW + 20, CH + 20),
                       yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        cv.pack(fill="both", expand=True)
        vsb.config(command=cv.yview)
        hsb.config(command=cv.xview)

        state = {
            "current_page": 0,
            "pages": [0]
        }

        # Preview Controls (Zoom & Navigation)
        ctrl_frame = tk.Frame(ws)
        ctrl_frame.pack(fill="x", pady=4)

        def prev_page():
            if state["current_page"] > 0:
                draw_page_preview(state["current_page"] - 1)

        def next_page():
            if state["current_page"] < len(state["pages"]) - 1:
                draw_page_preview(state["current_page"] + 1)

        # Page buttons hidden for single-page preview

        def zoom_in():
            if self.preview_scale < 2.0:
                self.preview_scale = round(min(2.0, self.preview_scale + 0.1), 1)
                draw_page_preview(state["current_page"])

        def zoom_out():
            if self.preview_scale > 0.25:
                self.preview_scale = round(max(0.25, self.preview_scale - 0.1), 1)
                draw_page_preview(state["current_page"])

        def on_mouse_wheel(event):
            if getattr(event, "state", 0) & 1:  # Shift pressed -> Horizontal scroll
                if getattr(event, "delta", 0) > 0 or getattr(event, "num", 0) == 4:
                    cv.xview_scroll(-1, "units")
                else:
                    cv.xview_scroll(1, "units")
            else:
                if getattr(event, "delta", 0) > 0 or getattr(event, "num", 0) == 4:
                    zoom_in()
                else:
                    zoom_out()
            return "break"

        cv.bind("<MouseWheel>", on_mouse_wheel)
        cv.bind("<Shift-MouseWheel>", on_mouse_wheel)
        cv.bind("<Button-4>", on_mouse_wheel)
        cv.bind("<Button-5>", on_mouse_wheel)

        btn_zoom_out = tk.Button(ctrl_frame, text="-", font=("Arial", 9, "bold"), command=zoom_out)
        btn_zoom_out.pack(side="left", padx=5)

        zoom_lbl = tk.Label(ctrl_frame, text="100%", font=("Arial", 10), bg="white")
        zoom_lbl.pack(side="left", padx=5)

        btn_zoom_in = tk.Button(ctrl_frame, text="+", font=("Arial", 9, "bold"), command=zoom_in)
        btn_zoom_in.pack(side="left", padx=5)

        def draw_page_preview(page_idx):
            state["current_page"] = page_idx
            cv.delete("all")
            
            scale = self.preview_scale
            w_scaled = int(PAGE_W * scale)
            h_scaled = int(PAGE_H * scale)
            mm_px = scale * 2.8346

            cv.update_idletasks()
            win_w = cv.winfo_width()
            content_w = max(win_w, w_scaled + 80) if win_w > 50 else w_scaled + 80
            OX = max(40, (content_w - w_scaled) // 2)
            OY = 10
            total_w = OX * 2 + w_scaled
            total_h = h_scaled + 40
            
            cv.config(scrollregion=(0, 0, total_w, total_h))
            
            cv.create_rectangle(OX+4, OY+4, OX+w_scaled+4, OY+h_scaled+4, fill="#888888", outline="")
            cv.create_rectangle(OX, OY, OX+w_scaled, OY+h_scaled, fill="white", outline="#aaaaaa", width=1)

            def ppx(pt):  return OX + int(pt * scale)
            def ppy(pt):  return OY + int((PAGE_H - pt) * scale)
            def pcx():    return OX + w_scaled // 2
            def smm(v):   return int(v * mm_px)

            zoom_lbl.config(text=f"{int(scale / base_scale * 100)}%")

            # Draw the logo
            logo_cx = ppx(20*mm)
            logo_cy = ppy(PAGE_H - 16*mm)
            logo_r  = smm(9)
            if os.path.isfile(LOGO_PATH):
                try:
                    from PIL import Image as _PI, ImageTk as _ITk
                    _img = _PI.open(LOGO_PATH)
                    _img.thumbnail((logo_r*2, logo_r*2))
                    self._pp_logo = _ITk.PhotoImage(_img)
                    cv.create_image(logo_cx, logo_cy, image=self._pp_logo)
                except Exception:
                    cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                                   logo_cx+logo_r, logo_cy+logo_r, outline="#555")
            else:
                cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                               logo_cx+logo_r, logo_cy+logo_r, outline="#555")

            # Header Text (match the generated PDF layout)
            header_cx = PAGE_W * 0.60
            header_cx1 = PAGE_W * 0.85
            cv.create_text(ppx(header_cx), ppy(PAGE_H - 15*mm), text=CLINIC_NAME,
                           font=("Times New Roman", max(10, smm(6)), "bold"), fill="#1a1a1a", anchor="center")
            cv.create_text(ppx(header_cx), ppy(PAGE_H - 20*mm), text=CLINIC_ADDRESS,
                           font=("Arial", max(8, smm(3))), fill="#333", anchor="center")
            cv.create_text(ppx(header_cx1), ppy(PAGE_H - 25*mm), text=CLINIC_PHONE,
                           font=("Arial", max(8, smm(3))), fill="#333", anchor="center")
            cv.create_text(ppx(header_cx1), ppy(PAGE_H - 30*mm), text=CLINIC_RESI,
                           font=("Arial", max(8, smm(3))), fill="#333", anchor="center")

            # Horizontal rule under header
            rule_y   = ppy(PAGE_H - 37*mm)
            marg_px  = ppx(1*mm)
            right_px = ppx(PAGE_W - 1*mm)
            cv.create_line(marg_px, rule_y, right_px, rule_y, fill="#333", width=1)
            cv.create_line(ppx(1*mm + 80), ppy(PAGE_H - 32*mm - 470), ppx(PAGE_W - 1*mm), ppy(PAGE_H - 32*mm - 470), fill="#333", width=1)
            
            #-------------Date & Time--------------------------------------
            date_box_x0 = ppx(PAGE_W - 30 * mm)
            date_box_y0 = ppy(175 * mm)
            # label inside box (top-left)
            cv.create_text(date_box_x0 + smm(1.5), date_box_y0 + smm(2), anchor="nw", text="Date :", font=("Helvetica", max(8, smm(3))))

            #-------------Doctor's Notes--------------------------------------
            note_divider_x = ppx(60 * mm)
            note_width = ppx(PAGE_W - 5 * mm) - note_divider_x
            note_height = smm(5)
            note_y = ppy(70 * mm)
            cv.create_rectangle(note_divider_x, note_y, note_divider_x + note_width, note_y + note_height,
                                outline="#555", width=1)
            cv.create_text(note_divider_x + note_width / 2, note_y + note_height / 2,
                           text="Doctor's Notes", font=("Helvetica", max(7, smm(3))), fill="#000")

            # Divider + writing-area box
            div_x = marg_px + smm(55)
            bot_y = ppy(12*mm)
            cv.create_line(div_x, rule_y, div_x, bot_y, fill="#555", width=1)
            cv.create_rectangle(div_x, rule_y, right_px, bot_y, outline="#555", width=1)

            # Left column: Consultants / Visiting doctors
            lx  = marg_px + smm(3)
            ly  = rule_y + smm(3)
            lh  = smm(4.6)
            lhs = smm(5.2)

            cv.create_text(lx, ly, text="Consultants :",
                           font=("Arial", max(7, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            ly += lhs
            for doc in CONSULTANTS:
                cv.create_text(lx+smm(3), ly, text=f"{doc['name']},",
                               font=("Arial", max(5, smm(2.5))), fill="#1a1a1a", anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=f"{doc['role']}",
                               font=("Arial", max(5, smm(2))), fill="#333",    anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=f"{doc['reg']}",
                               font=("Arial", max(5, smm(2))), fill="#555",    anchor="nw"); ly += lh + smm(1.5)

            ly += smm(2)
            cv.create_text(lx, ly, text="Visiting :",
                           font=("Arial", max(7, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            ly += lhs
            for doc in VISITING_DOCTORS:
                cv.create_text(lx+smm(3), ly, text=f"{doc['name']},",
                               font=("Arial", max(5, smm(2.5))), fill="#1a1a1a", anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=f"{doc['role']}",
                               font=("Arial", max(5, smm(2))), fill="#333",    anchor="nw"); ly += lh
                cv.create_text(lx+smm(3), ly, text=f"{doc['reg']}",
                               font=("Arial", max(5, smm(2))), fill="#555",    anchor="nw"); ly += lh + smm(1.5)

            # Footer
            cv.create_text(pcx(), ppy(6*mm), text=CLINIC_HOURS,
                           font=("Arial", max(6, smm(2.5))), fill="#555")

        draw_page_preview(0)

        # Buttons
        btn_bar = tk.Frame(ws)
        btn_bar.pack(pady=8)

        def _open_pdf(path):
            open_file(path)

        def _generate_pdf():
            filepath = os.path.join(SCRIPT_DIR, "Anupam_Dental_Clinic_Plain_Prescription.pdf")
            try:
                generate_pdf(filepath)
                docx_path = convert_pdf_to_docx(filepath)
                messagebox.showinfo("Done", f"Prescription saved successfully:\n📄 PDF: {filepath}\n📝 Word: {docx_path}")
                _open_pdf(filepath)
            except Exception as exc:
                messagebox.showerror("Error", f"PDF generation failed:\n{exc}")

        def _generate_word():
            filepath = os.path.join(SCRIPT_DIR, "Anupam_Dental_Clinic_Plain_Prescription.docx")
            try:
                tmp_pdf = os.path.join(SCRIPT_DIR, "_plain_prescription_temp.pdf")
                generate_pdf(tmp_pdf)
                docx_path = convert_pdf_to_docx(tmp_pdf, filepath)
                messagebox.showinfo("Done", f"Word Prescription saved:\n{docx_path}")
                open_file(docx_path)
            except Exception as exc:
                messagebox.showerror("Error", f"Word generation failed:\n{exc}")

        def _print_now():
            tmp = os.path.join(SCRIPT_DIR, "_plain_prescription_temp.pdf")
            try:
                generate_pdf(tmp)
                convert_pdf_to_docx(tmp)
                _open_pdf(tmp)
            except Exception as exc:
                messagebox.showerror("Error", f"Could not open PDF:\n{exc}")

        tk.Button(btn_bar, text="📄  Generate PDF", font=("Arial", 11), width=16,
                  bg="#1565C0", fg="white", command=_generate_pdf).grid(row=0, column=0, padx=6)
        tk.Button(btn_bar, text="📝  Word", font=("Arial", 11), width=12,
                  bg="#6A1B9A", fg="white", command=_generate_word).grid(row=0, column=2, padx=6)
        tk.Button(btn_bar, text="🖨  Open / Print", font=("Arial", 11), width=16,
                  bg="#2E7D32", fg="white", command=_print_now).grid(row=0, column=1, padx=6)
        tk.Button(btn_bar, text="Close",            font=("Arial", 11), width=10,
                  command=self.close).grid(row=0, column=3, padx=6)

    def Plain_Priscription_PDF_only(self):
        SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
        pdf_path = os.path.join(SCRIPT_DIR, "Anupam_Dental_Clinic_Plain_Prescription.pdf")
        
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import HexColor
        from reportlab.lib.utils import ImageReader
        
        A4_W, A4_H = A4
        PAGE_W, PAGE_H = A4_H / 2, A4_W
        LOGO_PATH = os.path.join(SCRIPT_DIR, "Dental_logo.png")
        CLINIC_NAME = "ANUPAM DENTAL CLINIC"
        CLINIC_ADDRESS = "West Gate Vaikom - 686141"
        CLINIC_PHONE, CLINIC_RESI = self.get_clinic_phone_settings()
        db_consultants, db_visiting = self.get_doctors_from_db()
        CONSULTANTS = list(db_consultants)
        VISITING_DOCTORS = list(db_visiting)

        def draw_logo(c: canvas.Canvas, cx: float, cy: float, r: float):
            if os.path.isfile(LOGO_PATH):
                img = ImageReader(LOGO_PATH)
                iw, ih = img.getSize()
                box = 2 * r
                scale = min(box / iw, box / ih)
                w, h = iw * scale, ih * scale
                c.drawImage(img, cx - w / 2, cy - h / 2, width=w, height=h, mask="auto", preserveAspectRatio=True)
                return

        def draw_pdf(c: canvas.Canvas):
            ink = HexColor("#1a1a1a")
            c.setFillColor(ink)
            c.setStrokeColor(ink)

            # Logo
            draw_logo(c, 20 * mm, PAGE_H - 16 * mm, 9 * mm)

            header_cx = PAGE_W * 0.60 
            header_cx1 = PAGE_W * 0.85

            c.setFont("Times-Bold", 20)
            c.drawCentredString(header_cx, PAGE_H - 15 * mm, CLINIC_NAME)

            c.setFont("Helvetica", 8.5)
            c.drawCentredString(header_cx, PAGE_H - 20 * mm, CLINIC_ADDRESS)

            c.setFont("Helvetica", 9)
            c.drawCentredString(header_cx1, PAGE_H - 25 * mm, CLINIC_PHONE)

            c.setFont("Helvetica", 9)
            c.drawCentredString(header_cx1, PAGE_H - 30 * mm, CLINIC_RESI)

            top_line_y = PAGE_H - 32 * mm
            margin = 1 * mm
            c.setLineWidth(0.8)
            c.line(margin, top_line_y, PAGE_W - margin, top_line_y) 
            c.line(margin + 80, top_line_y - 470, PAGE_W - margin, top_line_y - 470)

            divider_x = margin + 55 * mm
            bottom_line_y = 12 * mm
            c.setLineWidth(0.8)
            c.line(divider_x, top_line_y, divider_x, bottom_line_y)
            c.rect(divider_x, bottom_line_y, (PAGE_W - margin) - divider_x, top_line_y - bottom_line_y, fill=0, stroke=1)

            # Left column: Consultants / Visiting doctors
            lx = margin + 3 * mm
            ly = top_line_y - 4 * mm
            lh = 4.6 * mm

            for doc in CONSULTANTS:
                c.setFont("Helvetica-Bold", 7.5)
                c.drawString(lx, ly, f"{doc['name']},")
                ly -= lh
                c.setFont("Helvetica", 6.5)
                c.drawString(lx, ly, f"{doc['role']}")
                ly -= lh
                c.drawString(lx, ly, f"{doc['reg']}")
                ly -= (lh + 1.5 * mm)

            if VISITING_DOCTORS:
                ly -= 2 * mm
                c.setFont("Helvetica-Bold", 7)
                c.drawString(lx, ly, "VISITING DOCTORS:")
                ly -= lh
                for vdoc in VISITING_DOCTORS:
                    c.setFont("Helvetica-Bold", 7)
                    c.drawString(lx, ly, f"{vdoc['name']},")
                    ly -= lh
                    c.setFont("Helvetica", 6)
                    c.drawString(lx, ly, f"{vdoc['role']}")
                    ly -= lh
                    c.drawString(lx, ly, f"{vdoc['reg']}")
                    ly -= (lh + 1 * mm)

            # Date box
            c.setFont("Helvetica", 8)
            c.drawString(PAGE_W - 30 * mm, PAGE_H - 45 * mm, f"Date: {datetime.now().strftime('%d-%m-%Y')}")

            # Doctor's Notes box
            note_w = (PAGE_W - 5 * mm) - divider_x
            note_h = 6.5 * mm
            note_y = PAGE_H - 140 * mm
            c.rect(divider_x, note_y, note_w, note_h, fill=0, stroke=1)
            c.setFont("Helvetica-Bold", 8)
            c.drawCentredString(divider_x + note_w / 2, note_y + 2 * mm, "Doctor's Notes")

        c = canvas.Canvas(pdf_path, pagesize=(PAGE_W, PAGE_H))
        draw_pdf(c)
        c.save()
        open_file(pdf_path)
        return pdf_path
    def todayapp(self):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm as rl_mm
            from reportlab.pdfgen import canvas as rl_canvas
            from reportlab.lib.colors import HexColor
            from reportlab.lib.utils import ImageReader
            reportlab_ok = True
            PAGE_W, PAGE_H = A4          # points  (595 × 842)
            mm = rl_mm
        except ImportError:
            reportlab_ok = False
            PAGE_W, PAGE_H = 595.0, 842.0
            rl_mm = 2.8346
            mm = 2.8346
            HexColor = None
            ImageReader = None
            rl_canvas = None

        # ── Paths & clinic data ──────────────────────────────────────────
        SCRIPT_DIR     = os.path.dirname(os.path.abspath(__file__))
        LOGO_PATH      = os.path.join(SCRIPT_DIR, "Dental_logo.png")
        CLINIC_NAME    = "ANUPAM DENTAL CLINIC"
        CLINIC_ADDRESS = "West Gate Vaikom - 686141"
        CLINIC_PHONE, CLINIC_RESI = self.get_clinic_phone_settings()
        CLINIC_HOURS   = self.get_clinic_hours_setting()

        # ==============================================================
        # PDF generation (ReportLab)
        # ==============================================================
        def draw_logo(c, cx: float, cy: float, r: float):
            if os.path.isfile(LOGO_PATH) and ImageReader:
                img = ImageReader(LOGO_PATH)
                iw, ih = img.getSize()
                box = 2 * r
                scale = min(box / iw, box / ih)
                w, h = iw * scale, ih * scale
                c.drawImage(img, cx - w / 2, cy - h / 2, width=w, height=h, mask="auto", preserveAspectRatio=True)
                return
            
            # Fallback
            c.saveState()
            c.setDash(1, 2)
            c.setLineWidth(1)
            c.circle(cx, cy, r, stroke=1, fill=0)
            c.circle(cx, cy, r - 2 * rl_mm, stroke=1, fill=0)
            c.restoreState()
            c.saveState()
            c.translate(cx, cy)
            c.rotate(90)
            c.setFont("Helvetica-Bold", 5.5)
            c.drawCentredString(0, r - 5 * rl_mm, "ANUPAM")
            c.restoreState()
            c.setFont("Helvetica", 4.5)
            c.drawCentredString(cx, cy - 1.5 * rl_mm, "DENTAL")
            c.drawCentredString(cx, cy - 4.5 * rl_mm, "CLINIC")

        def draw_header(c: canvas.Canvas) -> float:
            """Draws the logo + clinic name/address/phone block and the rule
            beneath it. Returns the y coordinate of that rule (top_line_y)."""
            INK = HexColor("#1a1a1a")
            c.setFillColor(INK)
            c.setStrokeColor(INK)

            logo_cx, logo_cy, logo_r = 24 * mm, PAGE_H - 24 * mm, 13 * mm
            draw_logo(c, logo_cx, logo_cy, logo_r)

            header_cx = PAGE_W / 2 + 8 * mm

            c.setFont("Times-Bold", 24)
            c.drawCentredString(header_cx, PAGE_H - 20 * mm, CLINIC_NAME)

            c.setFont("Helvetica", 9)
            c.drawCentredString(header_cx, PAGE_H - 27 * mm, CLINIC_ADDRESS)

            right_x = PAGE_W - 18 * mm
            c.setFont("Helvetica", 9.5)
            c.drawRightString(right_x, PAGE_H - 33 * mm, CLINIC_PHONE)
            c.drawRightString(right_x, PAGE_H - 38 * mm, CLINIC_RESI)

            top_line_y = PAGE_H - 41 * mm
            c.setLineWidth(0.8)
            c.line(15 * mm, top_line_y, PAGE_W - 15 * mm, top_line_y)

            return top_line_y


        # ----------------------------------------------------------------------
        # Appointments table
        # ----------------------------------------------------------------------
        def draw_appointments(c: canvas.Canvas, top_line_y: float, doctor_name: str, appointments: list):
            """Draws the 'Appointments' title, the Doctor field, and the
            Time/PName/Purpose/Duration table below the header rule."""
            INK = HexColor("#1a1a1a") if HexColor else None
            if INK:
                c.setFillColor(INK)
                c.setStrokeColor(INK)
            else:
                c.setFillColorRGB(0.1, 0.1, 0.1)
                c.setStrokeColorRGB(0.1, 0.1, 0.1)

            left_margin = 15 * mm
            right_margin = PAGE_W - 15 * mm

            # ---------- section title ----------
            y = top_line_y - 10 * mm
            c.setFont("Helvetica-Bold", 13)
            c.drawCentredString(PAGE_W / 2, y, "Appointments")

            # ---------- Doctor field ----------
            y -= 10 * mm
            c.setFont("Helvetica-Bold", 9.5)
            c.drawString(left_margin, y + 1.5 * mm, "Doctor :")

            box_x0 = left_margin + 20 * mm
            box_x1 = left_margin + 95 * mm
            box_h = 6 * mm
            c.setLineWidth(0.6)
            c.rect(box_x0, y - 1 * mm, box_x1 - box_x0, box_h, stroke=1, fill=0)
            c.setFont("Helvetica", 9.5)
            c.drawString(box_x0 + 2 * mm, y + 1 * mm, doctor_name)

            # ---------- table column boundaries ----------
            col_time = (left_margin, left_margin + 27 * mm)
            col_name = (col_time[1], col_time[1] + 38 * mm)
            col_purpose = (col_name[1], col_name[1] + 78 * mm)
            col_duration = (col_purpose[1], right_margin)
            cols = [col_time, col_name, col_purpose, col_duration]
            headers = ["Time", "PName", "Purpose", "Duration"]

            table_top = y - 8 * mm
            header_h = 7 * mm
            row_h = 7 * mm

            # ---------- header row ----------
            c.setLineWidth(0.7)
            c.rect(left_margin, table_top - header_h, right_margin - left_margin, header_h, stroke=1, fill=0)
            for (x0, x1) in cols[:-1]:
                c.line(x1, table_top, x1, table_top - header_h)
            c.setFont("Helvetica-Bold", 9.5)
            for (x0, x1), label in zip(cols, headers):
                c.drawString(x0 + 2 * mm, table_top - header_h + 2 * mm, label)

            # ---------- data rows ----------
            c.setFont("Helvetica", 9)
            row_y = table_top - header_h
            for time_s, name_s, purpose_s, duration_s in appointments:
                c.setLineWidth(0.5)
                c.rect(left_margin, row_y - row_h, right_margin - left_margin, row_h, stroke=1, fill=0)
                for (x0, x1) in cols[:-1]:
                    c.line(x1, row_y, x1, row_y - row_h)

                text_y = row_y - row_h + 2 * mm
                c.drawString(col_time[0] + 2 * mm, text_y, time_s)
                c.drawString(col_name[0] + 2 * mm, text_y, name_s)
                c.drawString(col_purpose[0] + 2 * mm, text_y, purpose_s)
                c.drawRightString(col_duration[1] - 2 * mm, text_y, duration_s)

                row_y -= row_h

            return row_y  # bottom of the table, if the caller needs it


        # ----------------------------------------------------------------------
        # Footer
        # ----------------------------------------------------------------------
        def draw_footer(c: canvas.Canvas):
            INK = HexColor("#1a1a1a") if HexColor else None
            if INK:
                c.setFillColor(INK)
            else:
                c.setFillColorRGB(0.1, 0.1, 0.1)
            c.setFont("Helvetica", 8.5)
            c.drawCentredString(PAGE_W / 2 + 8 * mm, 14 * mm, CLINIC_HOURS)

        def generate_pdf(filepath: str):
            if not rl_canvas: return
            c = rl_canvas.Canvas(filepath, pagesize=(PAGE_W, PAGE_H))
            top_line_y = draw_header(c)
            
            # Fetch appointments for PDF
            conn = sqlite3.connect("dental.db")
            cursor = conn.cursor()
            today_str = datetime.now().strftime("%d-%m-%Y")
            cursor.execute("SELECT Time, Patient_Name, Notes FROM Appointments WHERE Date = ? ORDER BY Time", (today_str,))
            rows = cursor.fetchall()
            
            # Fetch doctor
            cursor.execute("SELECT First_Name, Last_Name, Qualification FROM Doctors LIMIT 1")
            doc_row = cursor.fetchone()
            conn.close()
            
            if doc_row:
                doc_name = f"Dr. {doc_row[0]} {doc_row[1]} {doc_row[2]}".strip()
            else:
                doc_name = "Dr. ANOOP KUMAR. B D S"
                
            appointments_list = [(r[0] if r[0] else "", r[1] if r[1] else "", r[2] if r[2] else "", "") for r in rows]
            
            draw_appointments(c, top_line_y, doc_name, appointments_list)
            draw_footer(c)
            c.showPage()
            c.save()

        # ==============================================================
        # Build Tkinter workspace
        # ==============================================================
        self.app.clear_workspace()
        self.app.workspace = tk.Frame(self.app.root, bd=3, relief="solid")
        self.app.workspace.pack(padx=10, pady=10, fill="both", expand=True)
        ws = self.app.workspace

        AVAIL_H = 500
        AVAIL_W = 500
        base_scale = min(AVAIL_H / PAGE_H, AVAIL_W / PAGE_W)
        self.preview_scale = base_scale
        CW      = int(PAGE_W * self.preview_scale)
        CH      = int(PAGE_H * self.preview_scale)

        outer = tk.Frame(ws, bg="#c0c0c0")
        outer.pack(fill="both", expand=True, padx=10, pady=(6, 0))

        vsb = tk.Scrollbar(outer, orient="vertical")
        vsb.pack(side="right", fill="y")
        hsb = tk.Scrollbar(outer, orient="horizontal")
        hsb.pack(side="bottom", fill="x")

        cv = tk.Canvas(outer, bg="#c0c0c0", width=CW+10, height=CH+1,
                       scrollregion=(0, 0, CW + 20, CH + 20),
                       yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        cv.pack(fill="both", expand=True)
        vsb.config(command=cv.yview)
        hsb.config(command=cv.xview)

        state = {
            "current_page": 0,
            "pages": [0]
        }

        # Preview Controls (Zoom & Navigation)
        ctrl_frame = tk.Frame(ws)
        ctrl_frame.pack(fill="x", pady=4)

        def prev_page():
            if state["current_page"] > 0:
                draw_page_preview(state["current_page"] - 1)

        def next_page():
            if state["current_page"] < len(state["pages"]) - 1:
                draw_page_preview(state["current_page"] + 1)

        # Page buttons hidden for single-page preview

        def zoom_in():
            if self.preview_scale < 2.0:
                self.preview_scale = round(min(2.0, self.preview_scale + 0.1), 1)
                draw_page_preview(state["current_page"])

        def zoom_out():
            if self.preview_scale > 0.25:
                self.preview_scale = round(max(0.25, self.preview_scale - 0.1), 1)
                draw_page_preview(state["current_page"])

        def on_mouse_wheel(event):
            if getattr(event, "delta", 0) > 0 or getattr(event, "num", 0) == 4:
                zoom_in()
            else:
                zoom_out()
            return "break"

        cv.bind("<MouseWheel>", on_mouse_wheel)
        cv.bind("<Button-4>", on_mouse_wheel)
        cv.bind("<Button-5>", on_mouse_wheel)

        btn_zoom_out = tk.Button(ctrl_frame, text="-", font=("Arial", 9, "bold"), command=zoom_out)
        btn_zoom_out.pack(side="left", padx=5)

        zoom_lbl = tk.Label(ctrl_frame, text="100%", font=("Arial", 10), bg="white")
        zoom_lbl.pack(side="left", padx=5)

        btn_zoom_in = tk.Button(ctrl_frame, text="+", font=("Arial", 9, "bold"), command=zoom_in)
        btn_zoom_in.pack(side="left", padx=5)

        btn_reg = tk.Button(ctrl_frame, text="Open Registration Workspace", font=("Arial", 9, "bold"), bg="#4CAF50", fg="white", command=lambda: self.app.registration())
        btn_reg.pack(side="right", padx=10)

        def draw_page_preview(page_idx):
            state["current_page"] = page_idx
            cv.delete("all")
            
            scale = self.preview_scale
            w_scaled = int(PAGE_W * scale)
            h_scaled = int(PAGE_H * scale)
            mm_px = scale * 2.8346

            cv.update_idletasks()
            win_w = cv.winfo_width()
            content_w = max(win_w, w_scaled + 80) if win_w > 50 else w_scaled + 80
            OX = max(40, (content_w - w_scaled) // 2)
            OY = 10
            total_w = OX * 2 + w_scaled
            total_h = h_scaled + 40
            
            cv.config(scrollregion=(0, 0, total_w, total_h))
            
            cv.create_rectangle(OX+4, OY+4, OX+w_scaled+4, OY+h_scaled+4, fill="#888888", outline="")
            cv.create_rectangle(OX, OY, OX+w_scaled, OY+h_scaled, fill="white", outline="#aaaaaa", width=1)

            def ppx(pt):  return OX + int(pt * scale)
            def ppy(pt):  return OY + int((PAGE_H - pt) * scale)
            def pcx():    return OX + w_scaled // 2
            def smm(v):   return int(v * mm_px)

            zoom_lbl.config(text=f"{int(scale / base_scale * 100)}%")

            # Draw the logo
            logo_cx = ppx(24*mm)
            logo_cy = ppy(PAGE_H - 24*mm)
            logo_r  = smm(13)
            if os.path.isfile(LOGO_PATH):
                try:
                    from PIL import Image as _PI, ImageTk as _ITk
                    _img = _PI.open(LOGO_PATH)
                    _img.thumbnail((logo_r*2, logo_r*2))
                    self._ta_logo = _ITk.PhotoImage(_img)
                    cv.create_image(logo_cx, logo_cy, image=self._ta_logo)
                except Exception:
                    cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                                   logo_cx+logo_r, logo_cy+logo_r, outline="#555")
            else:
                cv.create_oval(logo_cx-logo_r, logo_cy-logo_r,
                               logo_cx+logo_r, logo_cy+logo_r, outline="#555")

            # Header Text
            cv.create_text(pcx() + smm(8), ppy(PAGE_H - 20*mm), text=CLINIC_NAME,
                           font=("Times New Roman", max(10, smm(7.2)), "bold"), fill="#1a1a1a")
            cv.create_text(pcx() + smm(8), ppy(PAGE_H - 27*mm), text=CLINIC_ADDRESS,
                           font=("Arial", max(8, smm(2.7))), fill="#333")
            cv.create_text(ppx(PAGE_W - 18*mm), ppy(PAGE_H - 33*mm), text=f"{CLINIC_PHONE}",
                           font=("Arial", max(8, smm(2.8))), fill="#333", anchor="ne")
            cv.create_text(ppx(PAGE_W - 18*mm), ppy(PAGE_H - 38*mm), text=f"{CLINIC_RESI}",
                           font=("Arial", max(8, smm(2.8))), fill="#333", anchor="ne")

            top_line_y = PAGE_H - 43 * mm
            rule_y = ppy(top_line_y)
            cv.create_line(ppx(15 * mm), rule_y, ppx(PAGE_W - 15 * mm), rule_y, fill="#1a1a1a", width=1)
            
            # Fetch appointments for preview
            conn = sqlite3.connect("dental.db")
            cursor = conn.cursor()
            today_str = datetime.now().strftime("%d-%m-%Y")
            cursor.execute("SELECT Time, Patient_Name, Notes FROM Appointments WHERE Date = ? ORDER BY Time", (today_str,))
            rows = cursor.fetchall()
            
            # Fetch doctor
            cursor.execute("SELECT First_Name, Last_Name, Qualification FROM Doctors LIMIT 1")
            doc_row = cursor.fetchone()
            conn.close()
            
            if doc_row:
                doctor_name = f"Dr. {doc_row[0]} {doc_row[1]} {doc_row[2]}".strip()
            else:
                doctor_name = "Dr. ANOOP KUMAR. B D S"
                
            appointments_list = [(r[0] if r[0] else "", r[1] if r[1] else "", r[2] if r[2] else "", "") for r in rows]
            
            # Draw Appointments
            left_margin = 15 * mm
            right_margin = PAGE_W - 15 * mm
            
            y = top_line_y - 10 * mm
            cv.create_text(pcx(), ppy(y), text="Appointments", font=("Arial", max(10, smm(4)), "bold"), fill="#1a1a1a")
            
            y -= 10 * mm
            cv.create_text(ppx(left_margin), ppy(y + 1.5 * mm), text="Doctor :", font=("Arial", max(8, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            
            box_x0 = ppx(left_margin + 20 * mm)
            box_x1 = ppx(left_margin + 95 * mm)
            box_y0 = ppy(y + 5 * mm)
            box_y1 = ppy(y - 1 * mm)
            cv.create_rectangle(box_x0, box_y0, box_x1, box_y1, outline="#555", width=1)
            cv.create_text(box_x0 + smm(2), ppy(y + 1 * mm), text=doctor_name, font=("Arial", max(8, smm(3))), fill="#1a1a1a", anchor="nw")
            
            # Table boundaries
            col_time_x0 = ppx(left_margin)
            col_time_x1 = ppx(left_margin + 27 * mm)
            col_name_x1 = ppx(left_margin + 65 * mm)
            col_purpose_x1 = ppx(left_margin + 143 * mm)
            col_duration_x1 = ppx(right_margin)
            
            table_top = y - 8 * mm
            header_h = 7 * mm
            row_h = 7 * mm
            
            # Draw header row
            cv.create_rectangle(col_time_x0, ppy(table_top), col_duration_x1, ppy(table_top - header_h), outline="#555", width=1)
            cv.create_line(col_time_x1, ppy(table_top), col_time_x1, ppy(table_top - header_h), fill="#555", width=1)
            cv.create_line(col_name_x1, ppy(table_top), col_name_x1, ppy(table_top - header_h), fill="#555", width=1)
            cv.create_line(col_purpose_x1, ppy(table_top), col_purpose_x1, ppy(table_top - header_h), fill="#555", width=1)
            
            cv.create_text(col_time_x0 + smm(2), ppy(table_top - header_h + 2 * mm), text="Time", font=("Arial", max(8, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            cv.create_text(col_time_x1 + smm(2), ppy(table_top - header_h + 2 * mm), text="PName", font=("Arial", max(8, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            cv.create_text(col_name_x1 + smm(2), ppy(table_top - header_h + 2 * mm), text="Purpose", font=("Arial", max(8, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            cv.create_text(col_purpose_x1 + smm(2), ppy(table_top - header_h + 2 * mm), text="Duration", font=("Arial", max(8, smm(3)), "bold"), fill="#1a1a1a", anchor="nw")
            
            # Draw rows
            row_y = table_top - header_h
            for time_s, name_s, purpose_s, duration_s in appointments_list:
                cv.create_rectangle(col_time_x0, ppy(row_y), col_duration_x1, ppy(row_y - row_h), outline="#555", width=1)
                cv.create_line(col_time_x1, ppy(row_y), col_time_x1, ppy(row_y - row_h), fill="#555", width=1)
                cv.create_line(col_name_x1, ppy(row_y), col_name_x1, ppy(row_y - row_h), fill="#555", width=1)
                cv.create_line(col_purpose_x1, ppy(row_y), col_purpose_x1, ppy(row_y - row_h), fill="#555", width=1)
                
                text_y = row_y - row_h + 2 * mm
                cv.create_text(col_time_x0 + smm(2), ppy(text_y), text=time_s, font=("Arial", max(7, smm(2.7))), fill="#1a1a1a", anchor="nw")
                cv.create_text(col_time_x1 + smm(2), ppy(text_y), text=name_s, font=("Arial", max(7, smm(2.7))), fill="#1a1a1a", anchor="nw")
                cv.create_text(col_name_x1 + smm(2), ppy(text_y), text=purpose_s, font=("Arial", max(7, smm(2.7))), fill="#1a1a1a", anchor="nw")
                cv.create_text(col_duration_x1 - smm(2), ppy(text_y), text=duration_s, font=("Arial", max(7, smm(2.7))), fill="#1a1a1a", anchor="ne")
                row_y -= row_h
                
            # Footer
            cv.create_text(pcx(), ppy(6*mm), text=CLINIC_HOURS, font=("Arial", max(6, smm(2.5))), fill="#555")

        draw_page_preview(0)

        # Single zoom control bar defined above

        btn_bar = tk.Frame(ws)
        btn_bar.pack(pady=8)

        def _open_pdf(path):
            try:
                if hasattr(os, "startfile"):
                    os.startfile(path)
                elif sys.platform.startswith("linux"):
                    subprocess.Popen(["xdg-open", path])
                elif sys.platform == "darwin":
                    subprocess.Popen(["open", path])
                else:
                    subprocess.Popen(["cmd", "/c", "start", "", path])
            except Exception as exc:
                messagebox.showwarning("Open PDF", f"Could not open the PDF automatically.\n{exc}")

        def _generate_pdf():
            if not reportlab_ok:
                messagebox.showerror("Missing Library",
                    "ReportLab is required.\nRun:  pip install reportlab")
                return
            filepath = os.path.join(SCRIPT_DIR, "Anupam_Dental_Clinic_Appointments.pdf")
            try:
                generate_pdf(filepath)
                docx_path = convert_pdf_to_docx(filepath)
                messagebox.showinfo("Done", f"Appointments sheet saved:\n📄 PDF: {filepath}\n📝 Word: {docx_path}")
                _open_pdf(filepath)
            except Exception as exc:
                messagebox.showerror("Error", f"PDF generation failed:\n{exc}")

        def _print_now():
            if not reportlab_ok:
                messagebox.showerror("Missing Library",
                    "ReportLab is required.\nRun:  pip install reportlab")
                return
            tmp = os.path.join(SCRIPT_DIR, "_letterhead_temp.pdf")
            try:
                generate_pdf(tmp)
                _open_pdf(tmp)
            except Exception as exc:
                messagebox.showerror("Error", f"Could not open PDF:\n{exc}")

        tk.Button(btn_bar, text="📄  Generate PDF", font=("Arial", 11), width=16,
                  bg="#1565C0", fg="white", command=_generate_pdf).grid(row=0, column=0, padx=8)
        tk.Button(btn_bar, text="🖨  Open / Print", font=("Arial", 11), width=16,
                  bg="#2E7D32", fg="white", command=_print_now).grid(row=0, column=1, padx=8)
        tk.Button(btn_bar, text="Close",            font=("Arial", 11), width=10,
                  command=self.close).grid(row=0, column=2, padx=8)

    def close(self):
        if hasattr(self.app, "registration"):
            self.app.registration()
        else:
            from registration import Registration
            r = Registration(self.app)
            r.registration_workspace()

